import re
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
DEFAULT_MAP_FILE = BASE_DIR / "data" / "merged_map.xlsx"
DEFAULT_LOOKUP_FILE = BASE_DIR / "data" / "potno_yp_merge.xlsx"
DEFAULT_OUTPUT_FILE = BASE_DIR / "data" / "tma_map_replaced.xlsx"


@dataclass
class FormattingResult:
    """Summary of a run_formatting() call, for both CLI printing and app display."""
    lookup_sheets: list          # [(sheet_name, columns, row_count, used_bool), ...]
    mapping_count: int
    replaced_count: int
    unmatched: list              # [(sheet_name, cell_coordinate, yp_value), ...]
    output_file: Path
    workbook: object = field(repr=False)  # the saved openpyxl Workbook (post-replacement)


# Accepted spellings for the patient-number column in the lookup file,
# checked in order -- 'pat_no' is the original convention, 'pat_num' has
# shown up in newer copies of the lookup spreadsheet.
PAT_NO_COLUMN_ALIASES = ("pat_no", "pat_num")


def build_lookup(lookup_file):
    """Read every sheet of the YP -> patient-number lookup file into one dict.

    Sheets missing the required columns are skipped (and reported) rather
    than raising, since a lookup workbook may contain unrelated sheets.
    """
    lookup = {}
    sheets_info = []
    xl = pd.ExcelFile(lookup_file)
    for sheet in xl.sheet_names:
        df = xl.parse(sheet)
        df.columns = df.columns.str.strip().str.lower()
        pat_no_col = next((c for c in PAT_NO_COLUMN_ALIASES if c in df.columns), None)
        has_cols = pat_no_col is not None and "yp_num" in df.columns
        sheets_info.append((sheet, list(df.columns), len(df), has_cols))
        if not has_cols:
            continue
        for _, row in df.iterrows():
            yp = str(row["yp_num"]).strip()
            pat = str(row[pat_no_col]).strip()
            lookup[yp] = pat
    return lookup, sheets_info


def run_formatting(map_file=None, lookup_file=None, output_file=None):
    """Replace YP-numbers with patient numbers in the TMA block map.

    Reads every cell of `map_file`; any cell whose value starts with "YP"
    (case-insensitive) is looked up in `lookup_file` and replaced in place.
    All other cell content and formatting is preserved. The result is saved
    to `output_file`. Returns a FormattingResult with the same counts the
    CLI printout shows.
    """
    map_file = Path(map_file) if map_file else DEFAULT_MAP_FILE
    lookup_file = Path(lookup_file) if lookup_file else DEFAULT_LOOKUP_FILE
    output_file = Path(output_file) if output_file else DEFAULT_OUTPUT_FILE

    lookup, sheets_info = build_lookup(lookup_file)

    wb = load_workbook(map_file)

    replaced = 0
    unmatched = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows():
            for cell in row:
                val = str(cell.value).strip() if cell.value is not None else ""
                if re.match(r"^YP", val, re.IGNORECASE):
                    if val in lookup:
                        cell.value = lookup[val]
                        replaced += 1
                    else:
                        unmatched.append((sheet_name, cell.coordinate, val))

    output_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_file)

    return FormattingResult(
        lookup_sheets=sheets_info,
        mapping_count=len(lookup),
        replaced_count=replaced,
        unmatched=unmatched,
        output_file=output_file,
        workbook=wb,
    )


def _print_result(result):
    for sheet, columns, rows, used in result.lookup_sheets:
        if used:
            print(f"Sheet: '{sheet}' — columns: {columns} — rows: {rows}")
        else:
            print(f"Sheet: '{sheet}' — skipped (columns: {columns})")
    print(f"Loaded {result.mapping_count} YP → pat_no mappings")
    print(f"Done! Replaced {result.replaced_count} cells → saved to {result.output_file}")
    if result.unmatched:
        print(f"\nWarning: {len(result.unmatched)} YP-numbers had no match:")
        for sheet, coord, val in result.unmatched:
            print(f"  Sheet '{sheet}' cell {coord}: {val}")


if __name__ == "__main__":
    _print_result(run_formatting())
