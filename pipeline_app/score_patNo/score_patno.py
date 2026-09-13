import os
import re
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent
DEFAULT_DATA_DIR = BASE_DIR / "data"
DEFAULT_MAP_FILE_NAME = "tma_map_replaced.xlsx"
DEFAULT_OUTPUT_FILE = BASE_DIR / "patient_scores.xlsx"

# Filename suffixes flagging a score file whose scale differs from the
# shared 0-3 IHC intensity grade the other score files use -- see the
# Highest_Score note in run_score_patno()'s docstring.
NON_IHC_SCALE_SUFFIXES = ("_percentage", "_hscore")


@dataclass
class ScorePatNoResult:
    """Summary of a run_score_patno() call, for both CLI printing and app display."""
    score_labels: list
    score_files: list                 # file names used as score sources
    rows_before_filter: int
    rows_after_filter: int
    duplicate_patient_count: int      # patients with more than one core (Score_2, ...)
    sheet_name_fixes: list            # [(map_sheet, score_file_label, score_sheet), ...] — see _normalize_sheet_key
    per_score_tables: dict            # {label: DataFrame} — one sheet per score file
    summary_table: object             # DataFrame — All_Scores_Summary
    hscore_summary_table: object      # DataFrame or None — H_Scores_Summary (Patient_No + each "_hscore" column)
    output_file: Path


def read_all_cells(wb):
    """Return DataFrame with columns: sheet, row, col, value."""
    rows = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    rows.append((sheet_name, cell.row, cell.column, cell.value))
    return pd.DataFrame(rows, columns=["sheet", "row", "col", "value"])


def _normalize_sheet_key(name):
    """Collapse runs of whitespace/underscores into a single underscore, so
    e.g. a score file's "AN_VERR 1" and the map's "AN_VERR_1" are still
    treated as the same sheet during the position join. Without this, a
    single typo'd space in one source file silently drops every position on
    that sheet from every score file's join -- no error, no warning, just a
    quietly smaller "matched" count.
    """
    return re.sub(r"[\s_]+", "_", str(name).strip())


def _natural_key(s):
    """Zero-pad numeric parts so string sort matches numeric order."""
    return "".join(p.zfill(10) if p.isdigit() else p.lower()
                   for p in re.split(r"(\d+)", str(s)))


def flatten_duplicates(df, patient_col, value_cols):
    """
    Pivot duplicate patients into extra columns (Score_2, Score_3, …),
    append Highest_Score, and sort by patient number.
    """
    df = df.copy()
    df["_occ"] = df.groupby(patient_col).cumcount()

    frames = []
    for col in value_cols:
        pivoted = df.pivot_table(
            index=patient_col, columns="_occ", values=col, aggfunc="first"
        )
        pivoted.columns = [
            col if i == 0 else f"{col}_{i + 1}"
            for i, _ in enumerate(pivoted.columns)
        ]
        frames.append(pivoted)

    result = pd.concat(frames, axis=1).reset_index()

    score_cols = [c for c in result.columns if c != patient_col]
    result["Highest_Score"] = result[score_cols].max(axis=1)

    result = result.sort_values(
        patient_col, key=lambda s: s.map(_natural_key)
    ).reset_index(drop=True)

    return result


# --- Styling helpers ---
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", start_color="2F4F8F")
DATA_FONT = Font(name="Arial", size=10)
CENTER = Alignment(horizontal="center", vertical="center")
_THIN = Side(style="thin", color="CCCCCC")
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def write_sheet(wb, title, df):
    ws = wb.create_sheet(title=title[:31])
    headers = list(df.columns)
    for c_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c_idx, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER
    for r_idx, row_vals in enumerate(df.itertuples(index=False), 2):
        for c_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = DATA_FONT
            cell.alignment = CENTER
            cell.border = BORDER
    for col in ws.columns:
        max_len = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max_len + 4
    ws.freeze_panes = "A2"


def run_score_patno(data_dir=None, map_file=None, output_file=None):
    """Join TMA scoring files with the patient-number map by cell position.

    `data_dir` holds the patient map (`map_file`, default
    "tma_map_replaced.xlsx") plus any number of score `.xlsx` files — every
    other `.xlsx` file in `data_dir` is treated as a score source, named
    after its filename. Patient numbers "x" (missing core) / "i"
    (identifier) are excluded. Duplicate cores for the same patient are
    pivoted into Score/Score_2/... plus a Highest_Score column. Writes one
    sheet per score file plus an All_Scores_Summary sheet to `output_file`.

    A score file named with a "_percentage" suffix (e.g.
    "CD44_SFF_percentage.xlsx") is treated as a % positive tumor cells score
    (0-100), and one with an "_hscore" suffix (e.g. "CD44_SFF_hscore.xlsx")
    as a classic H-score (0-300) -- both are different scales from the 0-3
    IHC intensity grade the other score files use. Each still gets its own
    sheet/column like any other score file, but is excluded from the
    All_Scores_Summary sheet's cross-file Highest_Score column, since mixing
    scales into that max would make it meaningless (a 0-300 H-score would
    always dominate a 0-3 max). Add further suffixes to NON_IHC_SCALE_SUFFIXES
    below for any other non-0-3-scale score file.
    """
    data_dir = Path(data_dir) if data_dir else DEFAULT_DATA_DIR
    map_file_name = Path(map_file).name if map_file else DEFAULT_MAP_FILE_NAME
    map_path = data_dir / map_file_name
    output_file = Path(output_file) if output_file else DEFAULT_OUTPUT_FILE

    score_files = sorted([
        f for f in os.listdir(data_dir)
        if f.endswith(".xlsx") and not f.startswith("~$") and f != map_file_name
    ])

    # Load patient number map
    map_wb = load_workbook(map_path, data_only=True)
    map_df = read_all_cells(map_wb)
    map_df = map_df.rename(columns={"value": "Patient_No"})
    map_df["_sheet_key"] = map_df["sheet"].map(_normalize_sheet_key)
    map_sheet_names = set(map_df["sheet"])

    # Load scores from each file and join by position
    score_dfs = {}
    sheet_name_fixes = []
    for fname in score_files:
        label = os.path.splitext(fname)[0]
        score_wb = load_workbook(data_dir / fname, data_only=True)
        sc_df = read_all_cells(score_wb).rename(columns={"value": label})
        sc_df["_sheet_key"] = sc_df["sheet"].map(_normalize_sheet_key)
        score_dfs[label] = sc_df

        for score_sheet in sorted(set(sc_df["sheet"])):
            if score_sheet in map_sheet_names:
                continue
            key = _normalize_sheet_key(score_sheet)
            for map_sheet in map_sheet_names:
                if _normalize_sheet_key(map_sheet) == key:
                    sheet_name_fixes.append((map_sheet, label, score_sheet))

    # Merge all on position key (normalized sheet name + row + col) --
    # joining on _sheet_key rather than the raw "sheet" column is what
    # tolerates the kind of mismatch sheet_name_fixes reports above.
    combined = map_df.copy()
    for label, sc_df in score_dfs.items():
        combined = combined.merge(sc_df[["_sheet_key", "row", "col", label]],
                                  on=["_sheet_key", "row", "col"], how="left")

    # Convert score columns to numeric
    score_labels = list(score_dfs.keys())
    for col in score_labels:
        combined[col] = pd.to_numeric(combined[col], errors="coerce")

    rows_before_filter = len(combined)

    # Filter out missing ("x") and identifier ("i") patient numbers
    EXCLUDED = {"x", "i"}
    combined = combined[
        ~combined["Patient_No"].astype(str).str.strip().str.lower().isin(EXCLUDED)
    ]
    rows_after_filter = len(combined)

    duplicate_patient_count = int((combined.groupby("Patient_No").size() > 1).sum())

    wb_out = Workbook()
    wb_out.remove(wb_out.active)

    # One sheet per score file: Patient_No | Score | Score_2 (if duplicate) | Highest_Score
    per_score_tables = {}
    for label in score_labels:
        out_df = flatten_duplicates(combined[["Patient_No", label]], "Patient_No", [label])
        rename = {label: "Score"}
        for c in out_df.columns:
            if c.startswith(f"{label}_"):
                rename[c] = f"Score_{c[len(label) + 1:]}"
        out_df = out_df.rename(columns=rename)
        per_score_tables[label] = out_df
        write_sheet(wb_out, label, out_df)

    # Summary sheet: highest score per patient per score file
    summary_df = (
        combined[["Patient_No"] + score_labels]
        .groupby("Patient_No", as_index=False)
        .max()
        .sort_values("Patient_No", key=lambda s: s.map(_natural_key))
        .reset_index(drop=True)
    )
    # Cross-file Highest_Score only makes sense across labels on the same
    # measurement scale (0-3 IHC intensity). A label ending in one of
    # NON_IHC_SCALE_SUFFIXES (e.g. "_percentage" -> 0-100, "_hscore" -> 0-300)
    # is on a different scale instead -- folding it into this max would just
    # make Highest_Score echo that label's value rather than report a
    # meaningful combined IHC intensity score, so such labels are excluded
    # here (they still get their own sheet/column).
    ihc_scale_labels = [
        l for l in score_labels
        if not l.lower().endswith(NON_IHC_SCALE_SUFFIXES)
    ]
    summary_df["Highest_Score"] = summary_df[ihc_scale_labels].max(axis=1)
    write_sheet(wb_out, "All_Scores_Summary", summary_df)

    # H_Scores_Summary: a cross-marker export for just the classic H-score
    # (0-300) labels -- e.g. CD44_SFF_hscore and GATA6_hscore each get their
    # own column(s) here, one row per patient, so H-scores across markers can
    # be compared/exported without wading through the IHC-intensity and
    # percentage columns in All_Scores_Summary. Omitted entirely (both the
    # sheet and the returned table) when the cohort has no "_hscore" file yet.
    #
    # Built from per_score_tables rather than summary_df: summary_df already
    # collapsed every core down to a single per-patient max via
    # combined.groupby("Patient_No").max(), so a patient with two cores would
    # show only one (the higher) H-score here instead of both. per_score_tables
    # still has each core as Score / Score_2 / ... (see flatten_duplicates),
    # so that per-core detail is preserved per marker.
    hscore_labels = [l for l in score_labels if l.lower().endswith("_hscore")]
    hscore_summary_df = None
    if hscore_labels:
        for label in hscore_labels:
            tbl = per_score_tables[label].copy()
            rename = {}
            for c in tbl.columns:
                if c == "Patient_No":
                    continue
                elif c == "Score":
                    rename[c] = label
                elif c == "Highest_Score":
                    rename[c] = f"{label}_Highest"
                else:  # Score_2, Score_3, ...
                    rename[c] = f"{label}_{c.split('_', 1)[1]}"
            tbl = tbl.rename(columns=rename)
            hscore_summary_df = tbl if hscore_summary_df is None else hscore_summary_df.merge(tbl, on="Patient_No", how="outer")
        hscore_summary_df = hscore_summary_df.sort_values(
            "Patient_No", key=lambda s: s.map(_natural_key)
        ).reset_index(drop=True)
        write_sheet(wb_out, "H_Scores_Summary", hscore_summary_df)

    output_file.parent.mkdir(parents=True, exist_ok=True)
    wb_out.save(output_file)

    return ScorePatNoResult(
        score_labels=score_labels,
        score_files=score_files,
        rows_before_filter=rows_before_filter,
        rows_after_filter=rows_after_filter,
        duplicate_patient_count=duplicate_patient_count,
        sheet_name_fixes=sheet_name_fixes,
        per_score_tables=per_score_tables,
        summary_table=summary_df,
        hscore_summary_table=hscore_summary_df,
        output_file=output_file,
    )


def _print_result(result):
    print(f"Saved: {result.output_file}")
    sheets = result.score_labels + ["All_Scores_Summary"]
    if result.hscore_summary_table is not None:
        sheets.append("H_Scores_Summary")
    print(f"Sheets: {sheets}")
    if result.sheet_name_fixes:
        print(f"\nNote: {len(result.sheet_name_fixes)} sheet-name mismatch(es) tolerated "
              f"(matched anyway by normalizing spaces/underscores):")
        for map_sheet, label, score_sheet in result.sheet_name_fixes:
            print(f"  map sheet {map_sheet!r} <-> {label} sheet {score_sheet!r}")


if __name__ == "__main__":
    _print_result(run_score_patno())
