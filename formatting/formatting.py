import pandas as pd
from openpyxl import load_workbook
import re

# ── CONFIG ──────────────────────────────────────────────────────────────────
MAP_FILE    = "data/merged_map.xlsx"       # Excel file with TMA block map data
LOOKUP_FILE = "data/potno_yp_merge.xlsx"        # Excel file with pat_no and yp_num columns
OUTPUT_FILE = "data/tma_map_replaced.xlsx"
# ────────────────────────────────────────────────────────────────────────────

# Build YP → pat_no lookup dict from both sheets
lookup = {}
xl = pd.ExcelFile(LOOKUP_FILE)
for sheet in xl.sheet_names:
    df = xl.parse(sheet)
    df.columns = df.columns.str.strip().str.lower()
    if "pat_no" not in df.columns or "yp_num" not in df.columns:
        print(f"Sheet: '{sheet}' — skipped (columns: {list(df.columns)})")
        continue
    for _, row in df.iterrows():
        yp  = str(row["yp_num"]).strip()
        pat = str(row["pat_no"]).strip()
        lookup[yp] = pat
    print(f"Sheet: '{sheet}' — columns: {list(df.columns)} — rows: {len(df)}")
    

print(f"Loaded {len(lookup)} YP → pat_no mappings")

# Process TMA map — preserve formatting with openpyxl
wb = load_workbook(MAP_FILE)

replaced = 0
unmatched = []

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    for row in ws.iter_rows():
        for cell in row:
            val = str(cell.value).strip() if cell.value is not None else ""
            if re.match(r"^YP", val, re.IGNORECASE):
                if val in lookup:
                    cell.value = lookup[val]
                    replaced += 1
                else:
                    unmatched.append((sheet_name, cell.coordinate, val))

wb.save(OUTPUT_FILE)

print(f"Done! Replaced {replaced} cells → saved to {OUTPUT_FILE}")
if unmatched:
    print(f"\nWarning: {len(unmatched)} YP-numbers had no match:")
    for sheet, coord, val in unmatched:
        print(f"  Sheet '{sheet}' cell {coord}: {val}")