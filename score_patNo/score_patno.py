import os
import re
import pandas as pd
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DATA_DIR = "data"
MAP_FILE = os.path.join(DATA_DIR, "tma_map_replaced.xlsx")
OUTPUT_FILE = "patient_scores.xlsx"

SCORE_FILES = sorted([
    f for f in os.listdir(DATA_DIR)
    if f.endswith(".xlsx") and not f.startswith("~$") and f != "tma_map_replaced.xlsx"
])

def read_all_cells(wb):
    """Return DataFrame with columns: sheet, row, col, value."""
    rows = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    rows.append((sheet_name, cell.row, cell.column, cell.value))
    return pd.DataFrame(rows, columns=["sheet", "row", "col", "value"])

# Load patient number map
map_wb = load_workbook(MAP_FILE, data_only=True)
map_df = read_all_cells(map_wb)
map_df = map_df.rename(columns={"value": "Patient_No"})

# Load scores from each file and join by position
score_dfs = {}
for fname in SCORE_FILES:
    label = os.path.splitext(fname)[0]
    score_wb = load_workbook(os.path.join(DATA_DIR, fname), data_only=True)
    sc_df = read_all_cells(score_wb).rename(columns={"value": label})
    score_dfs[label] = sc_df

# Merge all on position key
combined = map_df.copy()
for label, sc_df in score_dfs.items():
    combined = combined.merge(sc_df[["sheet", "row", "col", label]],
                              on=["sheet", "row", "col"], how="left")

# Convert score columns to numeric
score_labels = list(score_dfs.keys())
for col in score_labels:
    combined[col] = pd.to_numeric(combined[col], errors="coerce")

# Filter out missing ("x") and identifier ("i") patient numbers
EXCLUDED = {"x", "i"}
combined = combined[
    ~combined["Patient_No"].astype(str).str.strip().str.lower().isin(EXCLUDED)
]

def _natural_key(s):
    """Zero-pad numeric parts so string sort matches numeric order."""
    return "".join(p.zfill(10) if p.isdigit() else p.lower()
                   for p in re.split(r"(\d+)", str(s)))

def flatten_duplicates(df, patient_col, value_cols):
    """
    Pivot duplicate patients into extra columns (Score_2, Score_3, …),
    append Highest_Score, and sort by patient number.
    """
    df = df.copy()
    df["_occ"] = df.groupby(patient_col).cumcount()

    frames = []
    for col in value_cols:
        pivoted = df.pivot_table(
            index=patient_col, columns="_occ", values=col, aggfunc="first"
        )
        pivoted.columns = [
            col if i == 0 else f"{col}_{i + 1}"
            for i, _ in enumerate(pivoted.columns)
        ]
        frames.append(pivoted)

    result = pd.concat(frames, axis=1).reset_index()

    score_cols = [c for c in result.columns if c != patient_col]
    result["Highest_Score"] = result[score_cols].max(axis=1)

    result = result.sort_values(
        patient_col, key=lambda s: s.map(_natural_key)
    ).reset_index(drop=True)

    return result

# --- Styling helpers ---
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", start_color="2F4F8F")
DATA_FONT = Font(name="Arial", size=10)
CENTER = Alignment(horizontal="center", vertical="center")
thin = Side(style="thin", color="CCCCCC")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

def write_sheet(wb, title, df):
    ws = wb.create_sheet(title=title[:31])
    headers = list(df.columns)
    for c_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c_idx, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER
    for r_idx, row_vals in enumerate(df.itertuples(index=False), 2):
        for c_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = DATA_FONT
            cell.alignment = CENTER
            cell.border = BORDER
    for col in ws.columns:
        max_len = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max_len + 4
    ws.freeze_panes = "A2"

wb_out = Workbook()
wb_out.remove(wb_out.active)

# One sheet per score file: Patient_No | Score | Score_2 (if duplicate) | Highest_Score
for label in score_labels:
    out_df = flatten_duplicates(combined[["Patient_No", label]], "Patient_No", [label])
    rename = {label: "Score"}
    for c in out_df.columns:
        if c.startswith(f"{label}_"):
            rename[c] = f"Score_{c[len(label) + 1:]}"
    out_df = out_df.rename(columns=rename)
    write_sheet(wb_out, label, out_df)

# Summary sheet: highest score per patient per score file
summary_df = (
    combined[["Patient_No"] + score_labels]
    .groupby("Patient_No", as_index=False)
    .max()
    .sort_values("Patient_No", key=lambda s: s.map(_natural_key))
    .reset_index(drop=True)
)
summary_df["Highest_Score"] = summary_df[score_labels].max(axis=1)
write_sheet(wb_out, "All_Scores_Summary", summary_df)

wb_out.save(OUTPUT_FILE)
print(f"Saved: {OUTPUT_FILE}")
print(f"Sheets: {[s.title for s in wb_out.worksheets]}")
